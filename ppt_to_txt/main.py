from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

# 사용할 PPTX 파일의 경로 (분석하고자 하는 pptx 파일의 파일 경로를 지정)
FILE = "/Users/dongrizheng/Documents/GitHub/report-automate-tool/ppt_to_txt/DXel 주간보고 20260511.pptx"

def walk_all_shapes_recursively(shapes):
    """
    모든 도형(Shape)들을 재귀적으로 순회하여 하나씩 반환합니다.
    파워포인트(PPTX)에서 도형(Shape)은 슬라이드에 존재하는 기본 객체이며, 
    도형은 그룹(묶음)을 가질 수 있는데, 그룹 내부의 도형도 다시 재귀적으로 탐색합니다.

    shapes: pptx.shapes.shapetree.SlideShapes 등에서 얻은 도형들의 리스트/컬렉션
    return: 모든 도형(Shape) 객체를 하나씩 yield
    """
    for shape in shapes:
        # 만약 이 shape가 그룹도형이라면 내부 도형들까지 재귀로 순회
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from walk_all_shapes_recursively(shape.shapes)
        else:
            yield shape

def extract_paragraphs_from_textframe(text_frame):
    """
    TextFrame(텍스트박스, 도형의 텍스트영역) 객체에서 
    각 문단 단위로 텍스트를 뽑아 리스트로 반환합니다.
    각 문단에는 여러 runs가 있을 수도 있고, 없을 경우 그냥 para.text를 사용합니다.

    text_frame : 도형(Shape)의 text_frame 속성
    return : ['문단1', '문단2', ...] 식의 텍스트 목록
    """
    paragraphs = []
    for para in text_frame.paragraphs:
        # runs가 있으면 run별로 합침, 없으면 para.text 사용
        merged_text = "".join(run.text for run in para.runs) if para.runs else para.text
        merged_text = " ".join(merged_text.split())  # 중간 불필요한 공백 정리
        # 필요 없는 텍스트(예시에서는 online) 거름
        if merged_text and merged_text.lower() not in ("(*)online", "online"):
            paragraphs.append(merged_text)
    return paragraphs

def get_merged_cell_text(cell):
    """
    표의 셀이 병합된 셀인지 확인해서, 병합된 셀(is_spanned)이면 빈 문자열 반환,
    일반 셀이면 그 셀의 텍스트(TextFrame에서 추출해서) 반환.
    
    cell: pptx.table._Cell
    return: str (셀의 텍스트)
    """
    if getattr(cell, "is_spanned", False):
        return ""
    return "\n".join(extract_paragraphs_from_textframe(cell.text_frame))

def extract_table_text_to_string(table):
    """
    표(Table) 객체 전체의 텍스트를 읽어서 사람이 읽기 쉽게 포맷팅된 문자열로 반환합니다.

    여러 상황(헤더 길이, 본문 내용 등)에 따라 아래 규칙으로 포맷합니다:
    - 헤더가 3개 이하이고 내용이 길면 (예: 큰 설명문 중심): 블록 형식
    - 헤더 이후, 각 행이 첫 컬럼만 값이 있으면: 리스트 형식
    - 그 외: key: value | key: value 식의 리스트
    
    table: pptx.table.Table
    return: str
    """
    rows = []
    for r in range(len(table.rows)):
        row = [get_merged_cell_text(table.cell(r, c)) for c in range(len(table.columns))]
        if any(row):  # 값이 비어있지 않은 행만 추가
            rows.append(row)
    if not rows:
        return ""

    head = rows[0]         # 헤더: 첫 번째 행
    body = rows[1:]        # 본문: 나머지 행들
    long_content = any(len(cell) > 50 or "\n" in cell for row in rows for cell in row)

    # [블록] 헤더가 3 이하 + 내용이 길고 + 본문 있음
    if len(head) <= 3 and long_content and body:
        blocks = []
        for idx, title in enumerate(head):
            vals = [row[idx] for row in body if idx < len(row) and row[idx]]
            if vals:
                blocks.append(f"### {title.replace(chr(10), ' ')}\n" + "\n".join(vals))
        return "\n\n".join(blocks)

    # [리스트] 헤더 이후 모든 행이 첫 컬럼만 값이 있으면
    if body and all(not any(row[1:]) for row in body):
        lines_ = [head[0]]
        for row in body:
            if row[0]:
                lines_.append(f"- {row[0]}")
        return "\n".join(lines_)

    # [Key:Value | ...] 그 외의 경우
    keys = [h.replace("\n", " ") or f"c{i}" for i, h in enumerate(head)]
    recs = []
    for row in body:
        pairs = [f"{k}: {v.replace(chr(10), ' ')}" for k, v in zip(keys, row) if v]
        if pairs:
            recs.append("- " + " | ".join(pairs))
    return "\n".join(recs) if recs else " | ".join(keys)

def extract_shape_text(shape):
    """
    도형(Shape)에서 텍스트를 추출합니다.
    이 도형이 표(table)인지, 텍스트프레임(text_frame)인지 자동으로 구분해서 적절히 추출합니다.

    shape: pptx.shapes.base.Shape
    return: str (추출된 텍스트/표 텍스트)
    """
    if shape.has_table:
        return extract_table_text_to_string(shape.table)
    if shape.has_text_frame:
        return "\n".join(extract_paragraphs_from_textframe(shape.text_frame))
    return ""

def extract_all_text_from_pptx(filepath):
    """
    PPTX 파일의 모든 슬라이드를 순서대로 읽어서, 
    각 슬라이드 내에 존재하는 모든 도형(Shape)에서 텍스트를 추출하여
    포맷팅된 하나의 문자열로 반환합니다.

    슬라이드 별로 텍스트를 section 별로 구분해서 내보냅니다.

    filepath: 분석할 pptx 파일의 경로 (str)
    return: str (슬라이드별로 구분된 모든 텍스트)
    """
    output_slide_texts = []
    for slide_index, slide in enumerate(Presentation(filepath).slides, 1):
        seen_texts = set()  # 중복 텍스트 방지
        shape_texts = []
        for shape in walk_all_shapes_recursively(slide.shapes):
            extracted = extract_shape_text(shape).strip()
            if extracted and extracted not in seen_texts:
                seen_texts.add(extracted)
                shape_texts.append(extracted)
        slide_text = f"## Slide {slide_index}\n\n" + "\n\n".join(shape_texts)
        output_slide_texts.append(slide_text)
    return "\n\n---\n\n".join(output_slide_texts) + "\n"

# 메인 실행부분: 지정된 파일에서 텍스트 추출해서 출력
if __name__ == "__main__":
    print(extract_all_text_from_pptx(FILE))