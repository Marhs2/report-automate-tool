import unittest
from copy import deepcopy
from io import BytesIO

from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn

from pptx_fill import MASTER_PATH, fill_weekly_pptx
from weekly_deck import from_legacy, next_week_label, section_lines, week_labels


class WeekLabelTests(unittest.TestCase):
    def test_done_span_and_following_weekday_span(self):
        done, nxt = week_labels(["2026-09-07", "2026-09-09", "2026-09-11"])
        self.assertEqual(done, "9/7~9/11")
        self.assertEqual(nxt, "9/14~9/18")

    def test_single_day_uses_the_following_calendar_week(self):
        done, nxt = week_labels(["2026-09-14"])
        self.assertEqual(done, "9/14")
        self.assertEqual(nxt, "9/21~9/25")

    def test_next_week_label_from_done_span(self):
        self.assertEqual(next_week_label("9/18", "2026.09.18"), "9/21~9/25")
        self.assertEqual(next_week_label("9/14~9/18", "2026.09.18"), "9/21~9/25")
        self.assertEqual(next_week_label("12/28~12/31", "2026.12.31"), "1/4~1/8")


class DeckTests(unittest.TestCase):
    def test_legacy_projects_become_done_and_next_sections(self):
        deck = from_legacy(
            {
                "projects": [
                    {
                        "projectName": "대한전선",
                        "completedTasks": ["외관검사기 셋팅"],
                        "inProgressTasks": ["태블릿 15대 발주"],
                        "issues": ["모니터 재연결"],
                        "nextWeekPlans": ["경광등 25대 설치"],
                    }
                ]
            },
            selected_dates=["2026-09-07", "2026-09-11"],
            member_name="정동일",
            team_name="Technical Center",
        )
        self.assertEqual(deck["schema"], "weekly-deck-v1")
        self.assertEqual(deck["author"], "정동일")
        self.assertEqual(
            deck["done"],
            [
                {
                    "title": "대한전선",
                    "items": ["외관검사기 셋팅", "태블릿 15대 발주", "모니터 재연결"],
                }
            ],
        )
        self.assertEqual(
            deck["next"],
            [{"title": "대한전선", "items": ["경광등 25대 설치"]}],
        )

    def test_keeps_notices_events_and_collab(self):
        deck = from_legacy(
            {
                "schema": "weekly-deck-v1",
                "done": [{"title": "A", "items": ["x"]}],
                "next": [],
                "director": "김동심",
                "notices": [{"title": "공지", "body": ["내용"]}],
                "month_events": [{"when": "9/1", "title": "회의"}],
                "next_month_events": [{"when": "10월", "title": "행사"}],
                "collab": {"supports": ["DC", "tc"], "done": ["펌웨어 지원"], "next": []},
            }
        )
        self.assertEqual(deck["director"], "김동심")
        self.assertEqual(deck["notices"], [{"title": "공지", "body": ["내용"]}])
        self.assertEqual(deck["month_events"], [{"when": "9/1", "title": "회의"}])
        self.assertEqual(deck["next_month_events"], [{"when": "10월", "title": "행사"}])
        self.assertEqual(deck["collab"]["supports"], ["DC", "TC"])
        self.assertEqual(deck["collab"]["done"], ["펌웨어 지원"])

    def test_schema_report_next_week_follows_selected_dates(self):
        deck = from_legacy(
            {
                "schema": "weekly-deck-v1",
                "done": [{"title": "A", "items": ["x"]}],
                "week_label_done": "9/14~9/18",
                "week_label_next": "9/14~9/18",
            },
            selected_dates=["2026-09-14", "2026-09-18"],
        )
        self.assertEqual(deck["week_label_next"], "9/21~9/25")

    def test_section_lines_match_sample_numbering(self):
        lines = section_lines(
            [
                {"title": "국책과제", "items": ["RCMS 등록", "전자계약 완료"]},
                {"title": "경영지원", "items": ["급여대장 작성"]},
            ]
        )
        self.assertEqual(
            lines,
            [
                "1. 국책과제",
                "   - RCMS 등록",
                "   - 전자계약 완료",
                "2. 경영지원",
                "   - 급여대장 작성",
            ],
        )


def _slide_blob(prs):
    texts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        texts.append(cell.text_frame.text)
    return "\n".join(texts)


def _status_cell_xml(prs, col=0):
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_table:
                continue
            header = shape.table.cell(0, 0).text_frame.text
            if "진행" in header:
                return shape.table.cell(1, col).text_frame._txBody.xml
    raise AssertionError("status table missing")


class FillPptxTests(unittest.TestCase):
    def test_filled_deck_uses_numbered_sections_and_keeps_white_text(self):
        self.assertTrue(MASTER_PATH.is_file())
        payload = fill_weekly_pptx(
            {
                "memberName": "정동일",
                "teamName": "Technical Center",
                "selectedDate": ["2026-09-07", "2026-09-11"],
                "report": from_legacy(
                    {
                        "projects": [
                            {
                                "projectName": "일일보고취합",
                                "completedTasks": ["제출 현황 달력 정리"],
                                "nextWeekPlans": ["PPT 칸 구조 수정"],
                            }
                        ],
                    },
                    selected_dates=["2026-09-07", "2026-09-11"],
                    member_name="정동일",
                    team_name="Technical Center",
                ),
            }
        )
        prs = Presentation(BytesIO(payload))
        blob = _slide_blob(prs)
        self.assertIn("정동일", blob)
        self.assertIn("1. 일일보고취합", blob)
        self.assertIn("제출 현황 달력 정리", blob)
        self.assertIn("PPT 칸 구조 수정", blob)
        self.assertNotIn("급여명세서 배포", blob)
        self.assertNotIn("[일일보고취합]", blob)
        self.assertNotIn("외 ", blob)
        self.assertNotIn("김래현", blob)
        self.assertNotIn("조기 지급", blob)
        self.assertNotIn("공지", blob)
        xml = _status_cell_xml(prs)
        self.assertIn('val="bg1"', xml)

    def test_empty_notices_drop_the_notice_slide(self):
        payload = fill_weekly_pptx(
            {
                "memberName": "정동일",
                "selectedDate": ["2026-09-07", "2026-09-11"],
                "report": {
                    "schema": "weekly-deck-v1",
                    "done": [{"title": "TYM", "items": ["점검"]}],
                    "next": [],
                    "notices": [],
                },
            }
        )
        prs = Presentation(BytesIO(payload))
        blob = _slide_blob(prs)
        self.assertNotIn("공지", blob)
        self.assertIn("점검", blob)

    def test_filled_notices_events_and_collab_stay_in_pptx(self):
        payload = fill_weekly_pptx(
            {
                "memberName": "정동일",
                "teamName": "Technical Center",
                "selectedDate": ["2026-09-07", "2026-09-11"],
                "report": {
                    "schema": "weekly-deck-v1",
                    "director": "김동심",
                    "done": [{"title": "TYM", "items": ["점검"]}],
                    "next": [],
                    "notices": [{"title": "보안규정", "body": ["문건 검토"]}],
                    "month_events": [{"when": "9/15", "title": "월보고"}],
                    "next_month_events": [{"when": "10/1", "title": "결산"}],
                    "collab": {
                        "supports": ["DC"],
                        "done": ["배터리 교체 지원"],
                        "next": ["모니터링 지원"],
                    },
                },
            }
        )
        prs = Presentation(BytesIO(payload))
        blob = _slide_blob(prs)
        self.assertIn("공지", blob)
        self.assertIn("보안규정", blob)
        self.assertIn("문건 검토", blob)
        self.assertIn("월보고", blob)
        self.assertIn("결산", blob)
        self.assertIn("센터 협업 현황", blob)
        self.assertIn("배터리 교체 지원", blob)
        self.assertIn("김동심", blob)
        self.assertIn("센터장", blob)

    def test_overflow_adds_another_status_slide_instead_of_cutting(self):
        items = [f"업무 {i}" for i in range(40)]
        payload = fill_weekly_pptx(
            {
                "memberName": "정동일",
                "teamName": "TC",
                "selectedDate": ["2026-09-07", "2026-09-11"],
                "report": {
                    "schema": "weekly-deck-v1",
                    "done": [{"title": "TYM", "items": items}],
                    "next": [],
                    "month_events": [],
                    "next_month_events": [],
                    "notices": [],
                },
            }
        )
        prs = Presentation(BytesIO(payload))
        self.assertGreaterEqual(len(prs.slides), 3)
        blob = _slide_blob(prs)
        self.assertIn("업무 0", blob)
        self.assertIn("업무 39", blob)
        self.assertNotIn("외 ", blob)
