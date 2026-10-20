from __future__ import annotations

import io
import math
import re
from copy import deepcopy
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

from weekly_deck import (
    COLLAB_CENTERS,
    collab_centers,
    collab_has_content,
    deck_layout,
    from_legacy,
    section_lines,
    week_labels,
)

MASTER_PATH = Path(__file__).resolve().parent / "templates" / "weekly_member_master.pptx"
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"
LINES_WITH_EVENTS = 16
LINES_WITHOUT_EVENTS = 28
# 마스터의 본문 글자는 9pt다. 28줄이 들어가도록 작게 잡혀 있어, 내용이 적으면 표 아래가 비고 글이 안 읽힌다.
# 내용이 표 높이에 들어가는 한 가장 큰 크기를 고른다(줄 바꿈까지 어림). 페이지 나눔은 그대로 9pt 기준이다.
BODY_SIZES_PT = (14, 13, 12, 11, 10)
LINE_SPACING = 1.3          # 마스터 본문 문단의 줄 간격(130%)
CELL_MARGIN_PT = 14         # 표 칸 안쪽 여백(좌우 합, 위아래 합) 어림
HEADING_GAP = 0.5           # 프로젝트 제목 위 여백, 글자 크기의 배수
HEADING = re.compile(r"^(\d+\.|\[[^\]]*\]$)")   # "1. 프로젝트" 또는 "[프로젝트]" 제목 줄
EMU_PER_PT = 12700


def _walk_shapes(shapes):
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _walk_shapes(shape.shapes)
        else:
            yield shape


def _char_em(ch: str) -> float:
    return 0.95 if ord(ch) >= 0x1100 else 0.55


def _visual_lines(lines, size_pt: float, width_pt: float) -> float:
    """줄 바꿈을 어림해 칸 안에서 차지하는 줄 수. 프로젝트 제목 위 여백도 줄로 셈다."""
    usable = max(1.0, width_pt - CELL_MARGIN_PT)
    total = 0.0
    for index, line in enumerate(lines):
        width = sum(_char_em(ch) for ch in line) * size_pt
        total += max(1, math.ceil(width / usable))
        if index and HEADING.match(line):
            total += HEADING_GAP / LINE_SPACING
    return total


def body_font_size(columns, width_pt: float, height_pt: float):
    """각 칸의 줄 목록이 모두 들어가는 가장 큰 본문 크기(pt). 어느 크기도 안 맞으면 None(마스터 그대로)."""
    room = height_pt - CELL_MARGIN_PT
    for size in BODY_SIZES_PT:
        if all(_visual_lines(lines, size, width_pt) * size * LINE_SPACING <= room for lines in columns):
            return size
    return None


def _style_run(rpr, *, size_pt=None, bold=None, scheme_color=None):
    if size_pt is not None:
        rpr.set("sz", str(int(size_pt * 100)))
    if bold is not None:
        rpr.set("b", "1" if bold else "0")
    if scheme_color is not None:
        for old in rpr.findall(qn("a:solidFill")):
            rpr.remove(old)
        fill = rpr.makeelement(qn("a:solidFill"), {})
        fill.append(fill.makeelement(qn("a:schemeClr"), {"val": scheme_color}))
        rpr.insert(0, fill)


def _style_cell(cell, *, size_pt=None, bold_headings=False, scheme_color=None):
    """쓴 문단의 글자 크기 · 색을 바꾸고, 프로젝트 제목 줄("1. …")은 굵게 하고 두 번째부터 위를 띄운다."""
    paragraphs = cell.text_frame._txBody.findall(qn("a:p"))
    headings = 0
    for paragraph in paragraphs:
        text = "".join(node.text or "" for node in paragraph.iter(qn("a:t")))
        heading = bool(bold_headings and HEADING.match(text))
        for rpr in list(paragraph.iter(qn("a:rPr"))) + list(paragraph.iter(qn("a:endParaRPr"))):
            _style_run(rpr, size_pt=size_pt, bold=True if heading else None, scheme_color=scheme_color)
        if heading:
            headings += 1
            if headings > 1 and size_pt:
                ppr = paragraph.find(qn("a:pPr"))
                if ppr is None:
                    ppr = paragraph.makeelement(qn("a:pPr"), {})
                    paragraph.insert(0, ppr)
                gap = ppr.makeelement(qn("a:spcBef"), {})
                gap.append(gap.makeelement(qn("a:spcPts"), {"val": str(int(size_pt * HEADING_GAP * 100))}))
                after = ppr.find(qn("a:lnSpc"))
                ppr.insert(list(ppr).index(after) + 1 if after is not None else 0, gap)


def _write_paragraphs(cell, lines):
    tx_body = cell.text_frame._txBody
    existing = tx_body.findall(qn("a:p"))
    if not existing:
        cell.text_frame.text = "\n".join(lines or [""])
        return
    template = deepcopy(existing[0])
    for node in existing:
        tx_body.remove(node)
    rows = lines or [""]
    for line in rows:
        paragraph = deepcopy(template)
        runs = paragraph.findall(qn("a:r"))
        if runs:
            first = runs[0]
            text_node = first.find(qn("a:t"))
            if text_node is None:
                text_node = first.makeelement(qn("a:t"), {})
                first.append(text_node)
            text_node.text = line
            text_node.set(XML_SPACE, "preserve")
            for extra in runs[1:]:
                paragraph.remove(extra)
        tx_body.append(paragraph)


def _delete_shape(shape):
    shape._element.getparent().remove(shape._element)


def _set_textbox(shape, value):
    paragraph = shape.text_frame.paragraphs[0]
    if paragraph.runs:
        paragraph.runs[0].text = value
        for extra in paragraph.runs[1:]:
            extra.text = ""
    else:
        paragraph.add_run().text = value


def _fill_cover(slide, deck):
    for shape in _walk_shapes(slide.shapes):
        if shape.has_text_frame and shape.text_frame.text.strip().startswith("Center"):
            _set_textbox(
                shape, f"Center : {deck.get('center') or deck.get('author') or ''}".rstrip()
            )
        if not shape.has_table:
            continue
        table = shape.table
        if len(table.rows) < 2 or len(table.columns) < 2:
            continue
        if "보고일자" not in table.cell(0, 0).text_frame.text:
            continue
        _write_paragraphs(table.cell(0, 1), [deck.get("report_date") or ""])
        _write_paragraphs(
            table.cell(1, 1),
            [deck.get("director") or deck.get("author") or ""],
        )
        # 마스터는 이 표 글자를 어두운 본문색(tx1)으로 두어, 짙은 표지 배경에서 거의 안 보인다.
        for row in range(len(table.rows)):
            for col in range(len(table.columns)):
                _style_cell(table.cell(row, col), scheme_color="bg1")


def _fill_notices(slide, notices):
    blocks = []
    for index, notice in enumerate(notices or [], start=1):
        title = str(notice.get("title") or "").strip()
        body = [str(line).strip() for line in notice.get("body") or [] if str(line).strip()]
        if not title and not body:
            continue
        blocks.append(f"{index}. {title}" if title else f"{index}.")
        blocks.extend(f"   - {line}" for line in body)
    if not blocks:
        blocks = ["없음"]
    for shape in _walk_shapes(slide.shapes):
        if not shape.has_table:
            continue
        table = shape.table
        if "공지" not in table.cell(0, 0).text_frame.text:
            continue
        body_row = 1 if len(table.rows) > 1 else 0
        _write_paragraphs(table.cell(body_row, 0), blocks)
        if len(table.columns) > 1:
            _write_paragraphs(table.cell(body_row, 1), [""])
        header = table.rows[0].height if body_row else 0
        size = body_font_size([blocks], table.columns[0].width / EMU_PER_PT, (shape.height - header) / EMU_PER_PT)
        _style_cell(table.cell(body_row, 0), size_pt=size, bold_headings=True)
        return True
    return True


def _fill_events(table, month_events, next_events):
    data_rows = list(range(1, len(table.rows)))
    for offset, row in enumerate(data_rows):
        month = (month_events or [])[offset] if offset < len(month_events or []) else {}
        nxt = (next_events or [])[offset] if offset < len(next_events or []) else {}
        _write_paragraphs(table.cell(row, 0), [str(month.get("when") or "")])
        _write_paragraphs(table.cell(row, 1), [str(month.get("title") or "")])
        if len(table.columns) > 3:
            _write_paragraphs(table.cell(row, 2), [str(nxt.get("when") or "")])
            _write_paragraphs(table.cell(row, 3), [str(nxt.get("title") or "")])


def _expand_status_table(slide, event_shape):
    status = None
    for shape in slide.shapes:
        if not shape.has_table:
            continue
        header = shape.table.cell(0, 0).text_frame.text
        if "진행" in header or "향후" in header:
            status = shape
            break
    if status is None:
        return
    top = event_shape.top
    extra = event_shape.height
    _delete_shape(event_shape)
    status.top = top
    status.height = status.height + extra


def paginate_sections(sections, max_lines):
    pages = []
    current = []
    used = 0
    for section in sections or []:
        title = str(section.get("title") or "").strip()
        items = [str(item).strip() for item in section.get("items") or [] if str(item).strip()]
        if not title and not items:
            continue
        remaining = items
        first = True
        while remaining or first:
            first = False
            room = max_lines - used - 1
            if room < 1:
                if current:
                    pages.append(current)
                current = []
                used = 0
                continue
            take = remaining[:room]
            remaining = remaining[room:]
            current.append({"title": title, "items": take})
            used += 1 + len(take)
            if remaining:
                pages.append(current)
                current = []
                used = 0
    if current:
        pages.append(current)
    return pages or [[]]


def duplicate_slide(prs, index):
    source = prs.slides[index]
    dest = prs.slides.add_slide(source.slide_layout)
    for shape in list(dest.shapes):
        shape._element.getparent().remove(shape._element)
    tree = dest.shapes._spTree
    for shape in source.shapes:
        tree.insert_element_before(deepcopy(shape._element), "p:extLst")
    return dest


COLLAB_BODY_PT = 10
COLLAB_LEAD_SHARE = 0.12    # '지원' 열 폭. 나머지는 진행 현황 · 향후일정이 반씩.


def _drop_ext(node):
    """복사한 열 · 행의 a16 id를 떼 둔다. 같은 id가 두 번 나오면 PowerPoint가 파일 복구를 묻는다."""
    for ext in node.findall(qn("a:extLst")):
        node.remove(ext)
    return node


def _center_cell(cell, *, bold=True):
    cell._tc.get_or_add_tcPr().set("anchor", "ctr")
    for paragraph in cell.text_frame._txBody.findall(qn("a:p")):
        ppr = paragraph.find(qn("a:pPr"))
        if ppr is None:
            ppr = paragraph.makeelement(qn("a:pPr"), {})
            paragraph.insert(0, ppr)
        ppr.set("algn", "ctr")
        for rpr in list(paragraph.iter(qn("a:rPr"))) + list(paragraph.iter(qn("a:endParaRPr"))):
            _style_run(rpr, bold=bold)


def _add_lead_column(frame, lead_share, weights):
    """2열 업무 현황 표 앞에 이름 열을 더하고 본문 행을 len(weights)개로 늘린다. 행 높이는 weights 비율."""
    tbl = frame.table._tbl
    grid = tbl.tblGrid
    cols = grid.findall(qn("a:gridCol"))
    total = sum(int(col.get("w")) for col in cols)
    grid.insert(0, _drop_ext(deepcopy(cols[0])))
    lead = int(total * lead_share)
    half = (total - lead) // 2
    for col, width in zip(grid.findall(qn("a:gridCol")), (lead, half, total - lead - half)):
        col.set("w", str(width))
    header, body = tbl.findall(qn("a:tr"))[:2]
    for row in (header, body):
        first = row.find(qn("a:tc"))
        first.addprevious(deepcopy(first))
    room = frame.height - int(header.get("h"))
    rows = [body]
    for _ in weights[1:]:
        rows[-1].addnext(_drop_ext(deepcopy(body)))
        rows.append(rows[-1].getnext())
    for row, weight in zip(rows, weights):
        row.set("h", str(room * weight // sum(weights)))


def _fill_collab_slide(slide, deck):
    """센터 협업 현황: 지원 | 진행 현황 | 향후일정 표, 센터 다섯 줄은 비어도 둔다.

    업무 현황 표(2열)를 복제해 열 하나 · 행 넷을 더한다. 색 · 점선 · 글꼴을 마스터와 맞추려고.
    """
    tables = [shape for shape in _walk_shapes(slide.shapes) if shape.has_table]
    frame = next(
        (
            shape
            for shape in tables
            if "진행" in shape.table.cell(0, 0).text_frame.text or "향후" in shape.table.cell(0, 0).text_frame.text
        ),
        None,
    )
    if frame is None:
        return
    top = min(shape.top for shape in tables)
    bottom = frame.top + frame.height
    # 표가 아래 '(*)Online' 주석과 겹치지 않게 그 위에서 끝낸다.
    for shape in _walk_shapes(slide.shapes):
        if shape.has_text_frame and "Online" in shape.text_frame.text and shape.top > top:
            bottom = min(bottom, shape.top - EMU_PER_PT * 6)
    for shape in tables:
        if shape is not frame:
            _delete_shape(shape)
    frame.top = top
    frame.height = bottom - top
    _add_lead_column(frame, COLLAB_LEAD_SHARE, [1] * len(COLLAB_CENTERS))

    table = frame.table
    _write_paragraphs(table.cell(0, 0), ["지원"])
    _write_paragraphs(table.cell(0, 1), [f"진행 현황 ({deck.get('week_label_done') or ''})"])
    _write_paragraphs(table.cell(0, 2), [f"향후일정 ({deck.get('week_label_next') or ''})"])
    centers = collab_centers(deck.get("collab"))
    for row, code in enumerate(COLLAB_CENTERS, start=1):
        _write_paragraphs(table.cell(row, 0), [code])
        _center_cell(table.cell(row, 0))
        for col, key in ((1, "done"), (2, "next")):
            lines = [f"{index}. {text}" for index, text in enumerate(centers[code][key], start=1)]
            _write_paragraphs(table.cell(row, col), lines)
            _style_cell(table.cell(row, col), size_pt=COLLAB_BODY_PT)


ROWS_LEAD_SHARE = 0.16      # 'rows' 형식의 '구분'(프로젝트) 열 폭


def paginate_rows(done_sections, next_sections, max_lines):
    """'rows' 형식: 프로젝트마다 한 행(진행 · 향후를 같은 제목끼리 묶음)을 슬라이드별로 나눈다.

    shortcut: 한 프로젝트가 max_lines보다 길면 자르지 않고 한 장에 둔다(글자가 작아짐). 넘치는 일이 잦으면 항목 단위로 나눌 것.
    """
    rows = {}
    for key, sections in (("done", done_sections), ("next", next_sections)):
        for section in sections or []:
            title = str(section.get("title") or "").strip()
            items = [str(item).strip() for item in section.get("items") or [] if str(item).strip()]
            if not title and not items:
                continue
            rows.setdefault(title, {"title": title, "done": [], "next": []})[key] += items
    pages, current, used = [], [], 0
    for row in rows.values():
        cost = max(len(row["done"]), len(row["next"]), 1) + 1
        if current and used + cost > max_lines:
            pages.append(current)
            current, used = [], 0
        current.append(row)
        used += cost
    if current:
        pages.append(current)
    return pages or [[]]


def _status_pages(deck, max_lines):
    """업무 현황 슬라이드 한 장씩 (done_sections, next_sections, rows). rows는 'rows' 형식일 때만."""
    if deck_layout(deck) == "rows":
        return [([], [], page) for page in paginate_rows(deck.get("done"), deck.get("next"), max_lines)]
    done_pages = paginate_sections(deck.get("done") or [], max_lines)
    next_pages = paginate_sections(deck.get("next") or [], max_lines)
    count = max(len(done_pages), len(next_pages))
    done_pages += [[]] * (count - len(done_pages))
    next_pages += [[]] * (count - len(next_pages))
    return [(done, nxt, None) for done, nxt in zip(done_pages, next_pages)]


def _fill_status_rows(frame, *, done_label, next_label, rows):
    """'rows' 형식: 구분 | 진행 현황 | 향후일정 표, 프로젝트가 한 행."""
    rows = rows or [{"title": "", "done": [], "next": []}]
    _add_lead_column(frame, ROWS_LEAD_SHARE, [max(len(r["done"]), len(r["next"]), 1) + 1 for r in rows])
    table = frame.table
    _write_paragraphs(table.cell(0, 0), ["구분"])
    _write_paragraphs(table.cell(0, 1), [f"진행 현황 ({done_label})"])
    _write_paragraphs(table.cell(0, 2), [f"향후일정 ({next_label})"])
    width_pt = min(column.width for column in list(table.columns)[1:]) / EMU_PER_PT
    sizes = []
    for index, row in enumerate(rows, start=1):
        _write_paragraphs(table.cell(index, 0), [row["title"]])
        _center_cell(table.cell(index, 0))
        columns = [[f"- {text}" for text in row[key]] for key in ("done", "next")]
        for col, lines in enumerate(columns, start=1):
            _write_paragraphs(table.cell(index, col), lines)
        sizes.append(body_font_size(columns, width_pt, table.rows[index].height / EMU_PER_PT))
    size = None if None in sizes else min(sizes)
    for index in range(1, len(rows) + 1):
        for col in (1, 2):
            _style_cell(table.cell(index, col), size_pt=size)


def _fill_status_pair(
    slide, *, done_label, next_label, done_sections, next_sections, events, layout="numbered", rows=None
):
    event_shape = None
    status_shape = None
    for shape in _walk_shapes(slide.shapes):
        if not shape.has_table:
            continue
        header = shape.table.cell(0, 0).text_frame.text
        if "금월" in header:
            event_shape = shape
        elif "진행" in header or "향후" in header:
            status_shape = shape
    month_events = [e for e in (events[0] or []) if str(e.get("when") or e.get("title") or "").strip()]
    next_events = [e for e in (events[1] or []) if str(e.get("when") or e.get("title") or "").strip()]
    has_events = bool(month_events or next_events)
    if event_shape is not None and has_events:
        _fill_events(event_shape.table, month_events, next_events)
    elif event_shape is not None:
        _expand_status_table(slide, event_shape)
        status_shape = None
        for shape in _walk_shapes(slide.shapes):
            if shape.has_table and (
                "진행" in shape.table.cell(0, 0).text_frame.text
                or "향후" in shape.table.cell(0, 0).text_frame.text
            ):
                status_shape = shape
    if status_shape is None:
        return
    if layout == "rows":
        _fill_status_rows(status_shape, done_label=done_label, next_label=next_label, rows=rows)
        return
    table = status_shape.table
    _write_paragraphs(table.cell(0, 0), [f"진행 현황 ({done_label})"])
    _write_paragraphs(table.cell(0, 1), [f"향후일정 ({next_label})"])
    done_lines, next_lines = section_lines(done_sections, layout), section_lines(next_sections, layout)
    _write_paragraphs(table.cell(1, 0), done_lines)
    _write_paragraphs(table.cell(1, 1), next_lines)
    width_pt = min(column.width for column in table.columns) / EMU_PER_PT
    height_pt = (status_shape.height - table.rows[0].height) / EMU_PER_PT
    size = body_font_size([done_lines, next_lines], width_pt, height_pt)
    for col in (0, 1):
        _style_cell(table.cell(1, col), size_pt=size, bold_headings=True)


def _notice_slide_index(prs):
    for index, slide in enumerate(prs.slides):
        for shape in slide.shapes:
            if shape.has_text_frame and "공지" in shape.text_frame.text:
                return index
            if shape.has_table and "공지" in shape.table.cell(0, 0).text_frame.text:
                return index
    return None


def _status_slide_index(prs):
    for index, slide in enumerate(prs.slides):
        for shape in slide.shapes:
            if shape.has_table and "진행" in shape.table.cell(0, 0).text_frame.text:
                return index
    return None


def _delete_slide(prs, index):
    sld_id_lst = prs.slides._sldIdLst
    children = list(sld_id_lst)
    sld_id_lst.remove(children[index])
    rId = children[index].get(qn("r:id"))
    if rId:
        try:
            prs.part.drop_rel(rId)
        except Exception:
            pass


def _member_deck(report: dict) -> dict:
    selected = report.get("selectedDate") or []
    deck = from_legacy(
        report.get("report") or report,
        selected_dates=selected,
        member_name=str(report.get("memberName") or ""),
        team_name=str(report.get("teamName") or ""),
    )
    if selected:
        done_label, next_label = week_labels(selected)
        if not deck.get("week_label_done"):
            deck["week_label_done"] = done_label
        deck["week_label_next"] = next_label
    return deck


def _retitle(slide, needle: str, title: str) -> None:
    for shape in _walk_shapes(slide.shapes):
        if shape.has_text_frame and needle in shape.text_frame.text:
            _set_textbox(shape, title)
            return


def _append_member(prs, report: dict, cover_index: int, notice_index, status_index) -> None:
    """한 사람의 주간보고를 표지 · 공지 · 업무 현황 슬라이드로 뒤에 붙인다. 템플릿 슬라이드는 복사만 한다."""
    deck = _member_deck(report)
    who = " · ".join(x for x in (str(report.get("teamName") or ""), str(report.get("memberName") or "")) if x)
    _fill_cover(duplicate_slide(prs, cover_index), deck)

    notices = [
        n
        for n in (deck.get("notices") or [])
        if str(n.get("title") or "").strip() or any(str(x).strip() for x in n.get("body") or [])
    ]
    if notice_index is not None and notices:
        slide = duplicate_slide(prs, notice_index)
        _fill_notices(slide, notices)
        _retitle(slide, "공지", f"공지사항 · {who}")
    if status_index is None:
        return

    month_events = deck.get("month_events") or []
    next_events = deck.get("next_month_events") or []
    has_events = any(str(item.get("when") or item.get("title") or "").strip() for item in month_events + next_events)
    max_lines = LINES_WITH_EVENTS if has_events else LINES_WITHOUT_EVENTS
    for page, (done, nxt, rows) in enumerate(_status_pages(deck, max_lines)):
        slide = duplicate_slide(prs, status_index)
        _retitle(slide, "주요 업무", f"주요 업무 현황 · {who}")
        _fill_status_pair(
            slide,
            done_label=deck.get("week_label_done") or "",
            next_label=deck.get("week_label_next") or "",
            done_sections=done,
            next_sections=nxt,
            events=(month_events, next_events) if page == 0 else ([], []),
            layout=deck_layout(deck),
            rows=rows,
        )

    if collab_has_content(deck.get("collab")):
        slide = duplicate_slide(prs, status_index)
        _retitle(slide, "주요 업무", f"센터 협업 현황 · {who}")
        _fill_collab_slide(slide, deck)


def fill_combined_pptx(reports: list[dict], *, report_date: str, scope: str) -> bytes:
    """여러 사람의 주간보고를 한 파일로 합친다. 맨 앞에 전체 표지, 그다음 사람마다 표지 · 공지 · 업무 현황.

    reports는 fill_weekly_pptx가 받는 모양의 목록이고, 이 순서대로 실린다(팀 → 사람).
    """
    if not MASTER_PATH.is_file():
        raise FileNotFoundError(f"PPT 마스터가 없습니다: {MASTER_PATH}")
    prs = Presentation(str(MASTER_PATH))
    template_count = len(prs.slides)
    notice_index = _notice_slide_index(prs)
    status_index = _status_slide_index(prs)

    _fill_cover(duplicate_slide(prs, 0), {"center": scope, "report_date": report_date, "director": "전체 취합"})
    for report in reports:
        _append_member(prs, report, 0, notice_index, status_index)

    for index in reversed(range(template_count)):
        _delete_slide(prs, index)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def fill_weekly_pptx(report: dict) -> bytes:
    if not MASTER_PATH.is_file():
        raise FileNotFoundError(f"PPT 마스터가 없습니다: {MASTER_PATH}")
    selected = report.get("selectedDate") or []
    deck = from_legacy(
        report.get("report") or report,
        selected_dates=selected,
        member_name=str(report.get("memberName") or ""),
        team_name=str(report.get("teamName") or ""),
    )
    if selected:
        done_label, next_label = week_labels(selected)
        if not deck.get("week_label_done"):
            deck["week_label_done"] = done_label
        deck["week_label_next"] = next_label
    prs = Presentation(str(MASTER_PATH))
    if prs.slides:
        _fill_cover(prs.slides[0], deck)

    notices = [
        n
        for n in (deck.get("notices") or [])
        if str(n.get("title") or "").strip()
        or any(str(x).strip() for x in n.get("body") or [])
    ]
    notice_index = _notice_slide_index(prs)
    if notice_index is not None and notices:
        _fill_notices(prs.slides[notice_index], notices)

    status_index = _status_slide_index(prs)
    if status_index is None:
        if not notices:
            notice_index = _notice_slide_index(prs)
            if notice_index is not None:
                _delete_slide(prs, notice_index)
        buf = io.BytesIO()
        prs.save(buf)
        return buf.getvalue()

    month_events = deck.get("month_events") or []
    next_events = deck.get("next_month_events") or []
    has_events = any(
        str(item.get("when") or item.get("title") or "").strip()
        for item in month_events + next_events
    )
    max_lines = LINES_WITH_EVENTS if has_events else LINES_WITHOUT_EVENTS
    pages = _status_pages(deck, max_lines)

    status_slides = [prs.slides[status_index]]
    for _ in range(1, len(pages)):
        status_slides.append(duplicate_slide(prs, status_index))
    for page, slide in enumerate(status_slides):
        _fill_status_pair(
            slide,
            done_label=deck.get("week_label_done") or "",
            next_label=deck.get("week_label_next") or "",
            done_sections=pages[page][0],
            next_sections=pages[page][1],
            layout=deck_layout(deck),
            rows=pages[page][2],
            events=(month_events, next_events) if page == 0 else ([], []),
        )

    if collab_has_content(deck.get("collab")):
        collab_slide = duplicate_slide(prs, status_index)
        _retitle(collab_slide, "주요 업무", "센터 협업 현황")
        _fill_collab_slide(collab_slide, deck)

    if not notices:
        notice_index = _notice_slide_index(prs)
        if notice_index is not None:
            _delete_slide(prs, notice_index)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
