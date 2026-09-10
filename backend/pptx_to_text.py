"""PPTX → plain text extraction (from ppt_to_txt/main.py)."""

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def walk_all_shapes_recursively(shapes):
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from walk_all_shapes_recursively(shape.shapes)
        else:
            yield shape


def extract_paragraphs_from_textframe(text_frame):
    paragraphs = []
    for para in text_frame.paragraphs:
        merged_text = (
            "".join(run.text for run in para.runs) if para.runs else para.text
        )
        merged_text = " ".join(merged_text.split())
        if merged_text and merged_text.lower() not in ("(*)online", "online"):
            paragraphs.append(merged_text)
    return paragraphs


def get_merged_cell_text(cell):
    if getattr(cell, "is_spanned", False):
        return ""
    return "\n".join(extract_paragraphs_from_textframe(cell.text_frame))


def extract_table_text_to_string(table):
    rows = []
    for r in range(len(table.rows)):
        row = [
            get_merged_cell_text(table.cell(r, c))
            for c in range(len(table.columns))
        ]
        if any(row):
            rows.append(row)
    if not rows:
        return ""

    head = rows[0]
    body = rows[1:]
    long_content = any(
        len(cell) > 50 or "\n" in cell for row in rows for cell in row
    )

    if len(head) <= 3 and long_content and body:
        blocks = []
        for idx, title in enumerate(head):
            vals = [row[idx] for row in body if idx < len(row) and row[idx]]
            if vals:
                blocks.append(
                    f"### {title.replace(chr(10), ' ')}\n" + "\n".join(vals)
                )
        return "\n\n".join(blocks)

    if body and all(not any(row[1:]) for row in body):
        lines_ = [head[0]]
        for row in body:
            if row[0]:
                lines_.append(f"- {row[0]}")
        return "\n".join(lines_)

    keys = [h.replace("\n", " ") or f"c{i}" for i, h in enumerate(head)]
    recs = []
    for row in body:
        pairs = [
            f"{k}: {v.replace(chr(10), ' ')}"
            for k, v in zip(keys, row)
            if v
        ]
        if pairs:
            recs.append("- " + " | ".join(pairs))
    return "\n".join(recs) if recs else " | ".join(keys)


def extract_shape_text(shape):
    if shape.has_table:
        return extract_table_text_to_string(shape.table)
    if shape.has_text_frame:
        return "\n".join(extract_paragraphs_from_textframe(shape.text_frame))
    return ""


def extract_all_text_from_pptx(source):
    """Extract text from a PPTX path or file-like object."""
    output_slide_texts = []
    for slide_index, slide in enumerate(Presentation(source).slides, 1):
        seen_texts = set()
        shape_texts = []
        for shape in walk_all_shapes_recursively(slide.shapes):
            extracted = extract_shape_text(shape).strip()
            if extracted and extracted not in seen_texts:
                seen_texts.add(extracted)
                shape_texts.append(extracted)
        slide_text = f"## Slide {slide_index}\n\n" + "\n\n".join(shape_texts)
        output_slide_texts.append(slide_text)
    return "\n\n---\n\n".join(output_slide_texts) + "\n"
